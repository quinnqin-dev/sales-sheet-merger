import asyncio
import copy
import json
import os
import shutil
import sqlite3
import tempfile
import threading
import uuid
from contextlib import asynccontextmanager
from datetime import date
from decimal import Decimal, getcontext
from pathlib import Path

from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from starlette.exceptions import HTTPException as StarletteHTTPException
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from .engine import (ILLEGAL_XML, MAX_ROWS, convert_row, dumps, duplicate_key,
                     inspect_sheet, is_blank, read_rows, sheet_names, validate_rule, wire)
from .i18n import (LOCALE, DomainError, Problem, descriptor, error_fields, localize_reason, resolve_locale, text)
from .models import DEFAULT_TEMPLATE, InspectRequest, Library, Rule, RunRequest, Template

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = Path(os.environ.get('SALES_MERGER_DATA_DIR', str(Path.home() / 'Library/Application Support/SalesSheetMerger')))
STATE_LOCK = threading.RLock()
RUN_LOCK = threading.Lock()
BATCHES = {}
JOBS = {}
WORK_DIR = None


def library_path():
    return DATA_DIR / 'library.json'


def save_library(lib):
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    temporary = DATA_DIR / 'library.json.new'
    temporary.write_text(lib.model_dump_json(indent=2), encoding='utf-8')
    temporary.replace(library_path())


def load_library():
    with STATE_LOCK:
        if library_path().exists():
            try:
                return Library.model_validate_json(library_path().read_text('utf-8'))
            except (ValueError, OSError):
                raise Problem(500, 'library_unreadable') from None
        return Library(template=DEFAULT_TEMPLATE, template_configured=False)


@asynccontextmanager
async def lifespan(app):
    global WORK_DIR
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    scratch = DATA_DIR / 'temporary'
    scratch.mkdir(exist_ok=True)
    # Avoid deleting another running instance's active data.
    import fcntl
    handle = (DATA_DIR / 'instance.lock').open('a')
    try:
        fcntl.flock(handle, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except BlockingIOError:
        handle.close()
        raise RuntimeError(text('instance_running')) from None
    for item in scratch.iterdir():
        if item.is_dir():
            shutil.rmtree(item)
        else:
            item.unlink()
    WORK_DIR = Path(tempfile.mkdtemp(prefix='session-', dir=scratch))
    BATCHES.clear()
    JOBS.clear()
    try:
        yield
    finally:
        # Finish the bounded background task before removing its files.
        await asyncio.to_thread(RUN_LOCK.acquire)
        RUN_LOCK.release()
        shutil.rmtree(WORK_DIR, ignore_errors=True)
        handle.close()


app = FastAPI(title='Sales Sheet Merger', lifespan=lifespan)


@app.middleware('http')
async def local_only(request: Request, call_next):
    token = LOCALE.set(resolve_locale(request.headers.get('accept-language')))
    try:
        return await local_request(request, call_next)
    finally:
        LOCALE.reset(token)


async def local_request(request, call_next):
    host = request.headers.get('host', '').split(':')[0]
    if host not in ('127.0.0.1', 'localhost', 'testserver'):
        return problem_response(403, 'local_only')
    origin = request.headers.get('origin')
    if origin and origin not in (f'http://{request.headers.get("host")}', 'http://127.0.0.1:5173', 'http://localhost:5173'):
        return problem_response(403, 'invalid_origin')
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'no-referrer'
    response.headers['Content-Language'] = LOCALE.get()
    return response


def problem_response(status, code, params=None):
    return JSONResponse({'detail': text(code, params), 'code': code, 'params': params or {}}, status,
                        headers={'Content-Language': LOCALE.get()})


@app.exception_handler(StarletteHTTPException)
async def http_error(request, exc):
    if isinstance(exc, Problem):
        return problem_response(exc.status_code, exc.code, exc.params)
    return JSONResponse({'detail': exc.detail}, exc.status_code)


@app.exception_handler(RequestValidationError)
async def validation_error(request, exc):
    return problem_response(422, 'invalid_config')


def batch_for(batch_id):
    if batch_id not in BATCHES:
        raise Problem(404, 'batch_missing')
    return BATCHES[batch_id]


def file_for(batch_id, file_id):
    batch = batch_for(batch_id)
    if file_id not in batch['files']:
        raise Problem(404, 'file_missing')
    return batch['files'][file_id]


def public_file(item):
    result = {k: v for k, v in item.items() if k != 'path'}
    if item.get('error_code'):
        result.update(error_fields(item['error_code'], item.get('error_params')))
    result['sheets'] = [{**s, **error_fields(s['error_code'], s.get('error_params'))}
                        if s.get('error_code') else s for s in item['sheets']]
    return result


def inspect_file(item, options):
    item['encoding'] = options.encoding
    item['header_rows'] = options.header_rows
    try:
        names = sheet_names(item['path'])
        if len(names) > 100:
            raise DomainError('too_many_sheets')
        sheets = []
        for name in names:
            header_row = options.header_rows.get(name, 1)
            try:
                sheets.append(inspect_sheet(item['path'], name, header_row, options.encoding))
            except Exception as exc:
                sheets.append({'name': name, 'header_row': header_row, **error_fields(*descriptor(exc))})
        item['sheets'] = sheets
        item.update(error_fields('all_sheets_invalid' if all(s.get('error') for s in sheets) else None))
    except Exception as exc:
        item.update(error_fields(*descriptor(exc)))
        item['sheets'] = []


@app.get('/api/health')
def health():
    return {'status': 'ok', 'app': 'sales-sheet-merger', 'max_files': 30, 'max_rows': MAX_ROWS}


@app.get('/api/config')
def get_config():
    return load_library()


@app.put('/api/template')
def put_template(template: Template):
    with STATE_LOCK:
        lib = load_library()
        lib.template = template
        lib.template_configured = True
        save_library(lib)
    return lib


@app.post('/api/template/reset')
def reset_template():
    with STATE_LOCK:
        lib = load_library()
        lib.template = DEFAULT_TEMPLATE.model_copy(deep=True)
        lib.template_configured = False
        save_library(lib)
    return lib


@app.post('/api/rules')
def put_rule(rule: Rule):
    with STATE_LOCK:
        lib = load_library()
        validate_rule_or_400(rule, lib.template, rule.headers)
        lib.rules = [r for r in lib.rules if r.id != rule.id] + [rule]
        if len(lib.rules) > 500:
            raise Problem(400, 'rule_limit')
        save_library(lib)
    return lib


@app.delete('/api/rules/{rule_id}')
def delete_rule(rule_id: str):
    with STATE_LOCK:
        lib = load_library()
        lib.rules = [r for r in lib.rules if r.id != rule_id]
        save_library(lib)
    return lib


@app.get('/api/config/export')
def export_config():
    return JSONResponse(load_library().model_dump(), headers={'Content-Disposition': 'attachment; filename="sales-rules.json"'})


@app.post('/api/config/import')
async def import_config(request: Request):
    body = await request.body()
    if len(body) > 2 * 1024 * 1024:
        raise Problem(413, 'library_size')
    try:
        lib = Library.model_validate_json(body)
    except ValidationError:
        raise Problem(400, 'library_format') from None
    with STATE_LOCK:
        save_library(lib)
    return lib


@app.post('/api/batches')
def create_batch():
    if len(BATCHES) >= 50:
        raise Problem(400, 'batch_limit')
    batch_id = uuid.uuid4().hex
    directory = WORK_DIR / batch_id
    directory.mkdir()
    BATCHES[batch_id] = {'id': batch_id, 'directory': directory, 'files': {}}
    return {'id': batch_id}


@app.get('/api/batches/{batch_id}')
def get_batch(batch_id: str):
    return {'id': batch_id, 'files': [public_file(f) for f in batch_for(batch_id)['files'].values()]}


@app.post('/api/batches/{batch_id}/files')
async def upload_files(batch_id: str, files: list[UploadFile] = File(...)):
    batch = batch_for(batch_id)
    # Reservation prevents concurrent uploads from bypassing the file-count limit.
    with STATE_LOCK:
        if len(batch['files']) + len(files) > 30:
            raise Problem(400, 'file_limit')
        items = []
        for uploaded in files:
            file_id = uuid.uuid4().hex
            name = Path(uploaded.filename or 'Untitled').name
            suffix = Path(name).suffix.lower()
            item = {'id': file_id, 'name': name, 'encoding': 'utf-8-sig', 'header_rows': {},
                    'sheets': [], **error_fields('reading_file'), 'size': 0,
                    'path': batch['directory'] / f'{file_id}{suffix}'}
            batch['files'][file_id] = item
            items.append((uploaded, item, suffix))
    for uploaded, item, suffix in items:
        try:
            if suffix not in ('.xlsx', '.csv'):
                item.update(error_fields('unsupported_file'))
                continue
            with item['path'].open('wb') as dest:
                while chunk := await uploaded.read(1024 * 1024):
                    item['size'] += len(chunk)
                    if item['size'] > 50 * 1024 * 1024:
                        raise DomainError('file_size')
                    dest.write(chunk)
            await asyncio.to_thread(inspect_file, item, InspectRequest())
        except Exception as exc:
            item.update(error_fields(*descriptor(exc)))
            item['path'].unlink(missing_ok=True)
        finally:
            await uploaded.close()
    return [public_file(f) for f in batch['files'].values()]


@app.post('/api/batches/{batch_id}/files/{file_id}/inspect')
def reinspect_file(batch_id: str, file_id: str, options: InspectRequest):
    item = file_for(batch_id, file_id)
    inspect_file(item, options)
    return public_file(item)


@app.delete('/api/batches/{batch_id}/files/{file_id}')
def remove_file(batch_id: str, file_id: str):
    if any(j['batch_id'] == batch_id and j['status'] == 'running' for j in JOBS.values()):
        raise Problem(409, 'remove_while_running')
    item = file_for(batch_id, file_id)
    item['path'].unlink(missing_ok=True)
    del batch_for(batch_id)['files'][file_id]
    return {'ok': True}


def validate_rule_or_400(rule, template, headers):
    try:
        validate_rule(rule, template, headers)
    except ValueError as exc:
        raise Problem(400, exc.code if isinstance(exc, DomainError) else 'incomplete_mapping', **(exc.params if isinstance(exc, DomainError) else {})) from None


def prepare(batch_id, request):
    batch = batch_for(batch_id)
    if any(f['error'] for f in batch['files'].values()):
        raise Problem(400, 'unreadable_files')
    selected_keys, prepared, total = set(), [], 0
    for selection in request.selections:
        key = (selection.file_id, selection.sheet)
        if key in selected_keys:
            raise Problem(400, 'duplicate_selection')
        selected_keys.add(key)
        item = file_for(batch_id, selection.file_id)
        sheet = next((s for s in item['sheets'] if s['name'] == selection.sheet), None)
        if not sheet or sheet.get('error'):
            raise Problem(400, 'sheet_unreadable')
        if sheet['header_row'] != selection.header_row:
            raise Problem(400, 'header_changed')
        validate_rule_or_400(selection.rule, request.template, sheet['headers'])
        total += sheet['row_count']
        prepared.append((copy.deepcopy(item), copy.deepcopy(sheet), selection.model_copy(deep=True)))
    if total > MAX_ROWS:
        raise Problem(400, 'total_row_limit')
    return prepared, total


@app.post('/api/batches/{batch_id}/preview')
def preview(batch_id: str, request: RunRequest):
    prepared, total = prepare(batch_id, request)
    rows = []
    for item, sheet, selection in prepared:
        for n, values, missing, raw in read_rows(item['path'], sheet['name'], item['encoding']):
            if n <= selection.header_row or is_blank(raw):
                continue
            result, reasons = convert_row(values, missing, selection.rule, request.template, sheet['headers'])
            rows.append({'file': item['name'], 'sheet': sheet['name'], 'row': n,
                         'values': result, 'reasons': reasons, 'raw': [wire(v) for v in raw]})
            if len(rows) >= 20:
                return {'rows': rows, 'total': total}
    return {'rows': rows, 'total': total}


def connect(job):
    db = sqlite3.connect(job['db'])
    db.row_factory = sqlite3.Row
    return db


def run_job(job, prepared, request):
    db = None
    try:
        locale_token = LOCALE.set(job['locale'])
        getcontext().prec = 50  # This background thread owns its decimal context.
        db = connect(job)
        db.executescript('''
            CREATE TABLE rows (id INTEGER PRIMARY KEY, kind TEXT, data TEXT, raw TEXT,
                reasons TEXT, file TEXT, sheet TEXT, original_row INTEGER, fingerprint TEXT);
            CREATE TABLE duplicates (fingerprint TEXT PRIMARY KEY, count INTEGER);
            CREATE INDEX rows_kind ON rows(kind, id);
        ''')
        sources = []
        numeric_indices = [(i, c) for i, c in enumerate(request.template.columns) if c.type == 'number']
        totals = {c.id: Decimal(0) for _, c in numeric_indices}
        for item, sheet, selection in prepared:
            source = {'file': item['name'], 'sheet': sheet['name'], 'read': 0, 'valid': 0, 'issues': 0, 'blank': 0}
            for n, values, missing, raw in read_rows(item['path'], sheet['name'], item['encoding']):
                if n <= selection.header_row:
                    continue
                if is_blank(raw):
                    source['blank'] += 1
                    continue
                result, reasons = convert_row(values, missing, selection.rule, request.template, sheet['headers'])
                source['read'] += 1
                job['processed'] += 1
                if job['processed'] > MAX_ROWS:
                    raise DomainError('total_row_limit')
                kind = 'issues' if reasons else 'valid'
                source[kind] += 1
                key = duplicate_key(result, request.template) if not reasons else None
                db.execute('INSERT INTO rows(kind,data,raw,reasons,file,sheet,original_row,fingerprint) VALUES (?,?,?,?,?,?,?,?)',
                           (kind, dumps(result), dumps(raw), dumps(reasons), item['name'], sheet['name'], n, key))
                if not reasons:
                    db.execute('INSERT INTO duplicates VALUES (?,1) ON CONFLICT(fingerprint) DO UPDATE SET count=count+1', (key,))
                    for i, col in numeric_indices:
                        if result[i] is not None:
                            totals[col.id] += Decimal(result[i])
                if job['processed'] % 1000 == 0:
                    db.commit()
            sources.append(source)
        db.commit()
        dup = db.execute('SELECT COALESCE(SUM(count),0), COALESCE(SUM(count-1),0), COUNT(*) FROM duplicates WHERE count>1').fetchone()
        job['summary'] = {'read': job['processed'], 'valid': sum(s['valid'] for s in sources),
                          'issues': sum(s['issues'] for s in sources), 'blank': sum(s['blank'] for s in sources),
                          'duplicate_rows': dup[0], 'duplicate_extra': dup[1], 'duplicate_groups': dup[2],
                          'totals': {k: wire(v) for k, v in totals.items()}, 'sources': sources}
        job['phase_code'] = 'phase_exporting'
        export_workbook(job, db, request.template)
        job['status'] = 'done'
        job['phase_code'] = 'phase_done'
    except Exception as exc:
        job['status'] = 'failed'
        job.update(error_fields(exc.code, exc.params) if isinstance(exc, DomainError) else error_fields('processing_failed'))
    finally:
        if db:
            db.close()
        RUN_LOCK.release()
        LOCALE.reset(locale_token)


def text_cell(ws, value, heading=False):
    if isinstance(value, str):
        value = ILLEGAL_XML.sub(lambda match: f'\\u{ord(match.group()):04x}', value)
    cell = WriteOnlyCell(ws, value=value)
    if isinstance(value, str):
        cell.data_type = 's'
    if heading:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='214C40')
    return cell


def append_text(ws, values, heading=False):
    ws.append([text_cell(ws, v, heading) for v in values])


def export_workbook(job, db, template):
    book = Workbook(write_only=True)
    details = book.create_sheet(text('sheet_sales'))
    details.freeze_panes = 'A2'
    for index in range(1, len(template.columns) + 1):
        details.column_dimensions[get_column_letter(index)].width = 20
    append_text(details, [c.name for c in template.columns], True)
    for row in db.execute('SELECT data FROM rows WHERE kind="valid" ORDER BY id'):
        cells = []
        for value, col in zip(json.loads(row['data']), template.columns):
            if value is not None and col.type == 'number':
                cell = WriteOnlyCell(details, Decimal(value))
                cell.number_format = '#,##0.############'
            elif value is not None and col.type == 'date':
                cell = WriteOnlyCell(details, date.fromisoformat(value))
                cell.number_format = 'yyyy-mm-dd'
            else:
                cell = text_cell(details, value)
            cells.append(cell)
        details.append(cells)
    details.auto_filter.ref = f'A1:{get_column_letter(len(template.columns))}{job["summary"]["valid"] + 1}'
    book.save(job['output'])


@app.post('/api/batches/{batch_id}/jobs')
def start_job(batch_id: str, request: RunRequest):
    prepared, total = prepare(batch_id, request)
    if not RUN_LOCK.acquire(blocking=False):
        raise Problem(409, 'merge_running')
    job_id = uuid.uuid4().hex
    job = {'id': job_id, 'batch_id': batch_id, 'status': 'running', 'phase_code': 'phase_converting', 'locale': LOCALE.get(),
           'processed': 0, 'total': total, 'summary': None, 'error': None,
           'template': request.template.model_dump(), 'db': WORK_DIR / f'{job_id}.sqlite',
           'output': WORK_DIR / f'{job_id}.xlsx'}
    JOBS[job_id] = job
    # The request object and all selected source metadata are snapshots.
    threading.Thread(target=run_job, args=(job, prepared, request.model_copy(deep=True)), daemon=True).start()
    return public_job(job)


def public_job(job):
    result = {k: v for k, v in job.items() if k not in ('db', 'output')}
    result['phase'] = text(job['phase_code'])
    if job.get('error_code'):
        result.update(error_fields(job['error_code'], job.get('error_params')))
    return result


def job_for(job_id):
    job = JOBS.get(job_id)
    if not job:
        raise Problem(404, 'job_missing')
    return job


@app.get('/api/jobs/{job_id}')
def get_job(job_id: str):
    return public_job(job_for(job_id))


@app.get('/api/jobs/{job_id}/rows')
def get_rows(job_id: str, kind: str = 'valid', page: int = 1, page_size: int = 50, duplicates_only: bool = False):
    job = job_for(job_id)
    if job['status'] != 'done':
        raise Problem(409, 'job_incomplete')
    if kind not in ('valid', 'issues') or page < 1 or not 1 <= page_size <= 200:
        raise Problem(400, 'invalid_pagination')
    db = connect(job)
    try:
        condition = 'r.kind=?' + (' AND d.count>1' if duplicates_only else '')
        total = db.execute(f'SELECT COUNT(*) FROM rows r LEFT JOIN duplicates d ON r.fingerprint=d.fingerprint WHERE {condition}', (kind,)).fetchone()[0]
        rows = db.execute(f'SELECT r.*,d.count AS duplicate_count FROM rows r LEFT JOIN duplicates d ON r.fingerprint=d.fingerprint WHERE {condition} ORDER BY r.id LIMIT ? OFFSET ?',
                          (kind, page_size, (page - 1) * page_size)).fetchall()
        return {'total': total, 'rows': [{'file': r['file'], 'sheet': r['sheet'], 'row': r['original_row'],
                    'values': json.loads(r['data']), 'raw': json.loads(r['raw']), 'reasons': [localize_reason(reason) for reason in json.loads(r['reasons'])],
                    'duplicate_group': r['fingerprint'][:12] if (r['duplicate_count'] or 0) > 1 else None} for r in rows]}
    finally:
        db.close()


@app.get('/api/jobs/{job_id}/export')
def download(job_id: str):
    job = job_for(job_id)
    if job['status'] != 'done':
        raise Problem(409, 'job_incomplete')
    return FileResponse(job['output'], filename=f'sales-merged-{date.today().isoformat()}.xlsx',
                        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


DIST = ROOT / 'frontend/dist'
if DIST.exists():
    app.mount('/', StaticFiles(directory=DIST, html=True), name='frontend')
