import io
import json
import re
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend import app as service
from backend.i18n import resolve_locale
from test_merger import client, csv_bytes, make_rule, request_for, run, upload


@pytest.mark.parametrize('header,expected', [(None,'en'),('fr-FR','en'),('zh-CN','zh-CN'),
    ('en-US,en;q=0.9','en'),('zh;q=0,en;q=1','en'),('en;q=0.5,zh-CN;q=0.9','zh-CN'),
    ('zh-CN;q=invalid','en')])
def test_language_negotiation(header, expected):
    assert resolve_locale(header) == expected


def test_default_english_errors_and_localized_validation(client):
    default = client.get('/api/batches/missing')
    assert default.status_code == 404
    assert default.json()['code'] == 'batch_missing'
    assert default.json()['detail'].startswith('Batch not found')
    chinese = client.get('/api/batches/missing',headers={'Accept-Language':'zh-CN'})
    assert chinese.json()['code'] == default.json()['code']
    assert '批次不存在' in chinese.json()['detail']
    assert chinese.headers['content-language'] == 'zh-CN'
    invalid = client.put('/api/template',json={},headers={'Accept-Language':'zh-CN'})
    assert invalid.status_code == 422 and invalid.json()['code'] == 'invalid_config'
    assert '配置格式' in invalid.json()['detail']
    assert client.get('/api/config').json()['template']['columns'][0]['name'] == 'Date'


def test_language_context_is_isolated_across_parallel_requests(client):
    def read(locale):
        return client.get('/api/batches/missing',headers={'Accept-Language':locale}).json()['detail']
    with ThreadPoolExecutor(max_workers=4) as executor:
        answers = list(executor.map(read,['en','zh-CN']*10))
    assert all(text.startswith('Batch') for text in answers[::2])
    assert all('批次不存在' in text for text in answers[1::2])


def test_stored_file_errors_retranslate_without_reimport(client):
    batch, files = upload(client,[('原始名称.csv',csv_bytes(['金额','金额'],[[1,2]]))])
    assert files[0]['sheets'][0]['error'].startswith('Headers contain')
    chinese = client.get(f'/api/batches/{batch}',headers={'Accept-Language':'zh-CN'}).json()['files'][0]
    assert chinese['name'] == '原始名称.csv'
    assert chinese['sheets'][0]['error_code'] == 'invalid_headers'
    assert '重复列名' in chinese['sheets'][0]['error']
    assert client.get(f'/api/batches/{batch}').json()['files'][0]['sheets'][0]['error'].startswith('Headers contain')


@pytest.mark.parametrize('locale,sheets', [('en',['Merged Data']),
                                         ('zh-CN',['合并明细'])])
def test_export_language_snapshot_and_business_values(client,locale,sheets):
    headers=['日期','订单号','收入']
    batch,files=upload(client,[('客户原始.csv',csv_bytes(headers,[['2026-10-27','000123','1.10'],['invalid','000124','2']]))])
    rule=make_rule(headers,{'date':'日期','order_id':'订单号','amount':'收入'})
    body=request_for(files,[rule])
    # Customer-authored names must not change when the language changes.
    body['template']['name']='客户标准模板'
    body['template']['columns'][1]['name']='客户订单编号'
    client.headers['Accept-Language']=locale
    job=run(client,batch,body)
    assert job['locale']==locale
    opposite='en' if locale=='zh-CN' else 'zh-CN'
    response=client.get(f'/api/jobs/{job["id"]}',headers={'Accept-Language':opposite}).json()
    assert response['locale']==locale and response['phase_code']=='phase_done'
    assert response['phase']==('Completed' if opposite=='en' else '处理完成')
    issues=client.get(f'/api/jobs/{job["id"]}/rows?kind=issues',headers={'Accept-Language':opposite}).json()
    reason=issues['rows'][0]['reasons'][0]
    assert reason['code']=='invalid_date' and reason['params']=={'format':'%Y-%m-%d'}
    assert ('Date does not match' in reason['reason']) == (opposite=='en')
    book=load_workbook(io.BytesIO(client.get(f'/api/jobs/{job["id"]}/export',headers={'Accept-Language':opposite}).content))
    assert book.sheetnames==sheets
    assert book[sheets[0]]['B1'].value=='客户订单编号'
    assert book[sheets[0]]['B2'].value=='000123'
    assert book[sheets[0]].max_row == 2
    assert issues['rows'][0]['file'] == '客户原始.csv'
    assert issues['rows'][0]['raw'][0] == 'invalid'
    book.close()


def test_existing_chinese_library_survives_restart(tmp_path,monkeypatch):
    monkeypatch.setattr(service,'DATA_DIR',tmp_path/'legacy')
    with TestClient(service.app) as session:
        library=session.get('/api/config').json()
        library['template']['name']='旧销售模板'
        library['template']['columns'][0]['name']='原始交易日期'
        rule=make_rule(['日期','金额'],{'date':'日期','amount':'金额'})
        library['rules']=[rule.model_dump()]
        assert session.post('/api/config/import',json=library).status_code==200
    with TestClient(service.app) as session:
        for locale in ('en','zh-CN'):
            assert session.get('/api/config',headers={'Accept-Language':locale}).json()==library
        batch,files=upload(session,[('中文.csv',csv_bytes(['日期','金额'],[['2026-10-27','3.10']]))])
        body=request_for(files,[rule]);body['template']=library['template']
        assert run(session,batch,body)['summary']['valid']==1


def test_catalog_keys_and_placeholders_match():
    directory=Path(__file__).resolve().parents[1]/'shared/locales'
    en=json.loads((directory/'en.json').read_text())
    zh=json.loads((directory/'zh-CN.json').read_text())
    assert en.keys()==zh.keys()
    for key in en:
        assert en[key] and zh[key]
        assert set(re.findall(r'\{\w+\}',en[key]))==set(re.findall(r'\{\w+\}',zh[key])),key
        assert not re.search(r'[\u4e00-\u9fff]',en[key]),key
