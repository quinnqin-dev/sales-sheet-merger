"""Real BIFF workbooks exercise the same import/preview/export path as XLSX/CSV."""
import io
from datetime import date, datetime

import pytest
import xlwt
from openpyxl import load_workbook

from backend.engine import read_rows
from test_merger import client, csv_bytes, excel, make_rule, request_for, run, upload


def legacy_excel(date_1904=False, title_row=False):
    book = xlwt.Workbook()
    book.set_dates_1904(date_1904)
    sheet = book.add_sheet('成绩明细')
    offset = int(title_row)
    if title_row:
        sheet.write(0, 0, '学生成绩')
    for column, name in enumerate(['日期', '编号', '金额']):
        sheet.write(offset, column, name)
    sheet.write(offset + 1, 0, datetime(2026, 10, 27), xlwt.easyxf(num_format_str='YYYY-MM-DD'))
    sheet.write(offset + 1, 1, 12, xlwt.easyxf(num_format_str='000000'))
    sheet.write(offset + 1, 2, 86.5)
    second = book.add_sheet('另一班')
    for column, value in enumerate(['日期', '编号', '金额']):
        second.write(0, column, value)
    for column, value in enumerate(['2026-10-28', '000013', 92]):
        second.write(1, column, value)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize('date_1904', [False, True])
def test_xls_mixed_merge_preview_export_and_rule_reuse(client, date_1904):
    headers = ['日期', '编号', '金额']
    batch, files = upload(client, [
        ('旧系统.XLS', legacy_excel(date_1904)),
        ('新系统.xlsx', excel(headers, [[date(2026, 10, 29), '000014', 1.25]])),
        ('平台.csv', csv_bytes(headers, [['2026-10-30', '000015', '2.25']])),
    ])
    assert not any(f['error'] for f in files)
    assert [s['name'] for s in files[0]['sheets']] == ['成绩明细', '另一班']
    assert files[0]['sheets'][0]['sample'][0]['values'] == ['2026-10-27T00:00:00', '000012', 86.5]
    rule = make_rule(headers, {'date': '日期', 'order_id': '编号', 'amount': '金额'})
    assert client.post('/api/rules', json=rule.model_dump()).status_code == 200
    saved_rule = client.get('/api/config').json()['rules'][0]
    body = request_for(files, [rule] * 3)
    body['selections'][0]['rule'] = saved_rule
    body['selections'].insert(1, dict(body['selections'][0], sheet='另一班'))
    preview = client.post(f'/api/batches/{batch}/preview', json=body).json()
    assert preview['rows'][0]['values'][0:2] == ['2026-10-27', '000012']
    assert not preview['rows'][0]['reasons']
    job = run(client, batch, body)
    assert job['summary']['valid'] == 4 and job['summary']['issues'] == 0
    assert job['summary']['totals']['amount'] == '182.00'
    with_book = load_workbook(io.BytesIO(client.get(f'/api/jobs/{job["id"]}/export').content))
    try:
        assert with_book.sheetnames == ['Sales Details']
        sheet = with_book.active
        assert sheet.max_row == 5
        assert sheet['A2'].value == datetime(2026, 10, 27)
        assert [sheet.cell(row, 2).value for row in range(2, 6)] == ['000012', '000013', '000014', '000015']
        assert sheet['G2'].value == 86.5
    finally:
        with_book.close()


def test_xls_header_correction_and_corrupt_file(client):
    batch, files = upload(client, [('标题.xls', legacy_excel(title_row=True)), ('损坏.xls', b'not a workbook')])
    assert files[0]['sheets'][0]['headers'] == ['学生成绩']
    assert files[1]['error'] and files[1]['sheets'] == []
    corrected = client.post(f'/api/batches/{batch}/files/{files[0]["id"]}/inspect',
                            json={'header_rows': {'成绩明细': 2}}).json()
    assert not corrected['error']
    assert corrected['sheets'][0]['header_row'] == 2
    assert corrected['sheets'][0]['sample'][0]['row'] == 3
    assert corrected['sheets'][0]['headers'] == ['日期', '编号', '金额']


def test_xls_cell_types_and_empty_formula_result(tmp_path):
    book = xlwt.Workbook()
    sheet = book.add_sheet('Types')
    for index, value in enumerate([123, True, '000123', xlwt.Formula('1+1')]):
        sheet.write(0, index, value)
    sheet.row(1).set_cell_error(0, '#DIV/0!')
    path = tmp_path / 'types.xls'
    book.save(str(path))
    rows = list(read_rows(path, 'Types'))
    assert rows[0][1] == [123, True, '000123', '']
    assert type(rows[0][1][0]) is int and type(rows[0][1][1]) is bool
    assert rows[1][1][0] == '#DIV/0!'


@pytest.mark.parametrize('locale', ['en', 'zh-CN'])
def test_unsupported_formats_describe_supported_types(client, locale):
    client.headers['Accept-Language'] = locale
    _, files = upload(client, [('unsupported.xlsm', b'not supported')])
    assert files[0]['error_code'] == 'unsupported_file'
    assert all(extension in files[0]['error'] for extension in ['.xlsx', '.xls', '.csv'])
