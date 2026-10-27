"""Streaming readers and the single conversion pipeline used by preview and jobs."""
import csv
import hashlib
import json
import math
import re
import unicodedata
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from itertools import zip_longest
from pathlib import Path

from openpyxl import load_workbook

from .models import Rule, Template
from .i18n import DomainError, localize_reason

MAX_ROWS = 100_000
MAX_COLUMNS = 1000
ILLEGAL_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def normalized_headers(headers):
    return [unicodedata.normalize('NFKC', str(h)).strip().casefold() for h in headers]


def header_signature(headers):
    return sorted(normalized_headers(headers))


def wire(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, 'f')
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def dumps(value):
    return json.dumps(value, ensure_ascii=False, default=wire)


def check_xlsx(path: Path):
    with zipfile.ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > 256 * 1024 * 1024:
            raise DomainError('xlsx_size')


def sheet_names(path):
    if path.suffix == '.csv':
        return ['CSV']
    check_xlsx(path)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()


def read_rows(path: Path, sheet: str, encoding='utf-8-sig'):
    """Yield (original row, values, missing-formula-cache columns, raw values)."""
    if path.suffix == '.csv':
        with path.open('r', encoding=encoding, newline='') as handle:
            sample = handle.read(8192)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            for n, values in enumerate(csv.reader(handle, dialect), 1):
                if len(values) > MAX_COLUMNS:
                    raise DomainError('column_limit')
                yield n, values, set(), values
        return
    values_book = load_workbook(path, read_only=True, data_only=True)
    formulas_book = None
    try:
        formulas_book = load_workbook(path, read_only=True, data_only=False)
        values_sheet, formulas_sheet = values_book[sheet], formulas_book[sheet]
        # Some exporters report A1:A1 despite having more data.
        values_sheet.reset_dimensions()
        formulas_sheet.reset_dimensions()
        for n, pair in enumerate(zip_longest(values_sheet.rows, formulas_sheet.rows, fillvalue=()), 1):
            value_cells, formula_cells = pair
            if max(len(value_cells), len(formula_cells)) > MAX_COLUMNS:
                raise DomainError('column_limit')
            values, raw, missing = [], [], set()
            for i, (cached, source) in enumerate(zip_longest(value_cells, formula_cells)):
                value = cached.value if cached else None
                original = source.value if source else None
                if source and source.data_type == 'f' and value is None:
                    missing.add(i)
                # Preserve identifiers displayed with simple zero-padding formats.
                if source and isinstance(value, (int, float)) and not isinstance(value, bool):
                    if re.fullmatch(r'0{2,}', source.number_format or '') and float(value).is_integer():
                        value = str(int(value)).zfill(len(source.number_format))
                values.append(value)
                raw.append(original if i in missing else value)
            yield n, values, missing, raw
    finally:
        values_book.close()
        if formulas_book:
            formulas_book.close()


def is_blank(values):
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in values)


def inspect_sheet(path, sheet, header_row, encoding):
    if not 1 <= header_row <= 1000:
        raise DomainError('header_range')
    headers, sample, count, blanks, physical = None, [], 0, 0, 0
    for n, values, missing, raw in read_rows(path, sheet, encoding):
        physical += 1
        if physical > 200_000:
            raise DomainError('physical_row_limit')
        if n == header_row:
            headers = [str(wire(v)).strip() if v is not None else '' for v in values]
            while headers and not headers[-1]:
                headers.pop()
            if missing:
                raise DomainError('header_formula')
        elif n > header_row:
            if is_blank(raw):
                blanks += 1
                continue
            count += 1
            if count > MAX_ROWS:
                raise DomainError('sheet_row_limit')
            if len(sample) < 20:
                sample.append({'row': n, 'values': [wire(v) for v in raw]})
    if not headers or not any(headers):
        raise DomainError('header_missing')
    normalized = normalized_headers(headers)
    if any(not h for h in normalized) or len(set(normalized)) != len(normalized):
        raise DomainError('invalid_headers')
    return {'name': sheet, 'header_row': header_row, 'headers': headers,
            'sample': sample, 'row_count': count, 'blank_count': blanks}


def validate_rule(rule: Rule, template: Template, headers):
    if header_signature(rule.headers) != header_signature(headers):
        raise DomainError('rule_headers')
    ids = {c.id for c in template.columns}
    if {m.target_id for m in rule.mappings} != ids:
        raise DomainError('template_changed')
    for m in rule.mappings:
        if m.mode == 'source' and (m.source_index is None or m.source_index >= len(rule.headers)):
            raise DomainError('incomplete_mapping')


def parse_number(value, style):
    if isinstance(value, bool) or isinstance(value, (datetime, date)):
        raise DomainError('invalid_number')
    text = str(value).strip()
    if style == 'comma' and ',' in text:
        if not re.fullmatch(r'[+-]?\d{1,3}(,\d{3})+(\.\d+)?', text):
            raise DomainError('comma_number')
        text = text.replace(',', '')
    elif style == 'european' and isinstance(value, str):
        if not re.fullmatch(r'[+-]?(\d+|\d{1,3}(\.\d{3})+)(,\d+)?', text):
            raise DomainError('european_number')
        text = text.replace('.', '').replace(',', '.')
    if not re.fullmatch(r'[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?', text):
        raise DomainError('number_symbols')
    try:
        result = Decimal(text)
    except InvalidOperation:
        raise DomainError('invalid_number') from None
    if not result.is_finite() or abs(result) > Decimal('1e15') or result.as_tuple().exponent < -12:
        raise DomainError('number_range')
    significant = ''.join(str(d) for d in result.as_tuple().digits).rstrip('0')
    if len(significant) > 15:
        raise DomainError('number_precision')
    return result


def convert_row(values, missing, rule, template, actual_headers):
    mappings = {m.target_id: m for m in rule.mappings}
    actual_positions = {name: i for i, name in enumerate(normalized_headers(actual_headers))}
    result, reasons = [], []
    for col in template.columns:
        m = mappings[col.id]
        try:
            if m.mode == 'source':
                source_name = normalized_headers([rule.headers[m.source_index]])[0]
                index = actual_positions[source_name]
                if index in missing:
                    raise DomainError('formula_cache')
                value = values[index] if index < len(values) else None
            elif m.mode == 'fixed':
                value = m.fixed
            else:
                value = None
            if isinstance(value, str) and m.trim:
                value = value.strip()
            if value is None or value == '':
                value = m.fill or None
            for replacement in m.replacements:
                if (str(wire(value)) if value is not None else '') == replacement.before:
                    value = replacement.after
                    break
            if value is None or value == '':
                if col.required:
                    raise DomainError('required_empty')
                result.append(None)
                continue
            if col.type == 'text':
                value = str(wire(value))
                if len(value) > 32767:
                    raise DomainError('text_length')
                if ILLEGAL_XML.search(value):
                    raise DomainError('text_control')
            elif col.type == 'number':
                value = parse_number(value, m.number_style)
            elif col.type == 'date':
                if isinstance(value, datetime):
                    value = value.date()
                elif not isinstance(value, date):
                    try:
                        value = datetime.strptime(str(value), m.date_format).date()
                    except ValueError:
                        raise DomainError('invalid_date', format=m.date_format) from None
            result.append(wire(value))
        except (ValueError, KeyError, IndexError) as exc:
            result.append(None)
            code = exc.code if isinstance(exc, DomainError) else 'invalid_value'
            params = exc.params if isinstance(exc, DomainError) else {}
            reasons.append(localize_reason({'column': col.name, 'code': code, 'params': params}))
    return result, reasons


def duplicate_key(values, template):
    canonical = [str(Decimal(v).normalize()) if c.type == 'number' and v is not None else v
                 for v, c in zip(values, template.columns)]
    return hashlib.sha256(dumps(canonical).encode()).hexdigest()
